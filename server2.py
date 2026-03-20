"""
AI Agent 백엔드 서버 (FastAPI) - 클라우드 배포 버전
- 브라우저에서 마이크 오디오를 받아 Whisper로 전사
- GPT-4o Function Calling으로 사용자 의도 파악
- Railway 등 클라우드 환경에 배포 가능

필요 패키지:
    pip install -r requirements.txt

로컬 실행:
    python server.py

Railway 배포:
    railway up
"""

import os
import io
import json
import datetime
import tempfile
from typing import Optional

from fastapi import FastAPI, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from openai import OpenAI

# ─────────────────────────────────────────
# 설정
# ─────────────────────────────────────────
API_KEY    = os.getenv("OPENAI_API_KEY", "")
OUTPUT_DIR = os.getenv("OUTPUT_DIR", ".")   # 회의록 저장 폴더
PORT       = int(os.getenv("PORT", 8000))   # Railway는 PORT 환경변수 사용

client = OpenAI(api_key=API_KEY)
app    = FastAPI(title="AI Agent Server")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─────────────────────────────────────────
# 요청/응답 모델
# ─────────────────────────────────────────
class ChatRequest(BaseModel):
    message: str
    context: Optional[str] = None

class ToolCallItem(BaseModel):
    name: str
    description: str
    status: str

class ChatResponse(BaseModel):
    reply: str
    tool_calls: list[ToolCallItem] = []

class RecordingStatus(BaseModel):
    is_recording: bool
    line_count: int
    minutes_path: str


# ─────────────────────────────────────────
# 녹음 세션 상태 (서버 메모리)
# ─────────────────────────────────────────
session = {
    "is_recording": False,
    "transcript":   [],       # 전사된 문장들
    "start_time":   None,
    "minutes_path": "",
}

HALLUCINATION_PATTERNS = [
    "mbc","kbs","sbs","뉴스","이덕영","앵커",
    "시청해주셔서 감사합니다","구독과 좋아요","구독","좋아요","알림설정",
    "감사합니다","안녕하세요","bye","thank you for watching",
    "자막 제공","번역 제공","subtitles","copyright",
]

def _is_hallucination(text: str) -> bool:
    t = text.lower().strip()
    if len(t) <= 2: return True
    return any(p in t for p in HALLUCINATION_PATTERNS)


# ─────────────────────────────────────────
# 회의록 생성
# ─────────────────────────────────────────
def _generate_minutes() -> str:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    transcript = session["transcript"]
    if not transcript:
        return ""

    start_time      = session["start_time"]
    transcript_text = "\n".join(transcript)

    prompt = f"""다음은 회의 전사 내용이야.
JSON 형식으로 회의록을 작성해줘:
{{
  "title": "회의 제목",
  "agenda": "안건 한 줄 요약",
  "discussions": [{{"topic": "주제", "content": "내용"}}],
  "decisions": ["결정 사항"],
  "action_items": [{{"task": "할 일", "owner": "담당자"}}]
}}

전사 내용:
{transcript_text}"""

    res  = client.chat.completions.create(
        model="gpt-4o",
        messages=[{"role": "user", "content": prompt}],
        response_format={"type": "json_object"},
    )
    data = json.loads(res.choices[0].message.content)

    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(3.0)

    def sf(run, size=11, bold=False, color=None):
        run.font.name = "맑은 고딕"; run.font.size = Pt(size); run.font.bold = bold
        if color: run.font.color.rgb = RGBColor(*color)
        rPr = run._r.get_or_add_rPr()
        rF  = OxmlElement('w:rFonts'); rF.set(qn('w:eastAsia'), "맑은 고딕"); rPr.insert(0, rF)

    def section_heading(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(12); p.paragraph_format.space_after = Pt(4)
        pPr = p._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
        bot = OxmlElement('w:bottom'); bot.set(qn('w:val'),'single'); bot.set(qn('w:sz'),'6')
        bot.set(qn('w:space'),'1'); bot.set(qn('w:color'),'1F497D')
        pBdr.append(bot); pPr.append(pBdr)
        r = p.add_run(text); sf(r, size=13, bold=True, color=(0x1F,0x49,0x7D))

    def bullet(text):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text); sf(r, size=10.5)

    tp = doc.add_paragraph(); tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = tp.add_run("회  의  록"); sf(r, size=22, bold=True, color=(0x1F,0x49,0x7D))
    sp = doc.add_paragraph(); sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(16)
    r2 = sp.add_run(data.get("title","회의")); sf(r2, size=12, color=(0x40,0x40,0x40))

    section_heading("1. 주요 논의 사항")
    for item in data.get("discussions", []):
        p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(4)
        r = p.add_run(f"▶ {item['topic']}"); sf(r, size=10.5, bold=True, color=(0x1F,0x49,0x7D))
        bullet(item["content"])

    section_heading("2. 결정 사항")
    for d in (data.get("decisions") or ["특별한 결정 사항 없음"]):
        bullet(d)

    section_heading("3. 후속 조치")
    for ai in (data.get("action_items") or [{"task":"없음","owner":""}]):
        bullet(f"{ai['task']}  {'— ' + ai['owner'] if ai.get('owner') else ''}")

    doc.add_paragraph().paragraph_format.space_after = Pt(20)
    fp = doc.add_paragraph(); fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r  = fp.add_run("※ 본 회의록은 STT + GPT-4o Agent가 자동 생성하였습니다.")
    sf(r, size=9, color=(0x80,0x80,0x80))

    fname = f"회의록_{start_time.strftime('%Y%m%d_%H%M')}.docx"
    out   = os.path.join(OUTPUT_DIR, fname)
    doc.save(out)
    return fname   # 파일명만 반환 (다운로드 URL용)


# ─────────────────────────────────────────
# Tool 정의
# ─────────────────────────────────────────
TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "start_meeting_recorder",
            "description": "실시간 회의 녹음을 시작한다. '회의 녹음', '회의록 만들어줘', 'STT' 요청에 사용.",
            "parameters": {
                "type": "object",
                "properties": {
                    "language": {"type": "string", "enum": ["ko","en","ja","zh"]}
                },
                "required": []
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "translate_text",
            "description": "텍스트를 지정 언어로 번역한다.",
            "parameters": {
                "type": "object",
                "properties": {
                    "text":        {"type": "string"},
                    "target_lang": {"type": "string"}
                },
                "required": ["text", "target_lang"]
            }
        }
    },
    {
        "type": "function",
        "function": {
            "name": "summarize_conversation",
            "description": "대화 내용을 요약한다.",
            "parameters": {
                "type": "object",
                "properties": {
                    "conversation": {"type": "string"}
                },
                "required": ["conversation"]
            }
        }
    },
]

def _run_tool(name: str, args: dict) -> str:
    if name == "start_meeting_recorder":
        if session["is_recording"]:
            return "이미 녹음이 진행 중입니다."
        return "RECORDING_START_REQUESTED"

    if name == "translate_text":
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role":"user","content":f"다음을 {args['target_lang']}로 번역해줘. 번역문만 출력.\n\n{args['text']}"}]
        )
        return res.choices[0].message.content.strip()

    if name == "summarize_conversation":
        res = client.chat.completions.create(
            model="gpt-4o",
            messages=[{"role":"user","content":f"다음 대화를 3~5줄로 한국어 요약:\n\n{args['conversation']}"}]
        )
        return res.choices[0].message.content.strip()

    return f"'{name}' 도구를 찾을 수 없습니다."


TOOL_DISPLAY = {
    "start_meeting_recorder": "회의 녹음 시작",
    "translate_text":         "텍스트 번역",
    "summarize_conversation": "대화 요약",
}


# ─────────────────────────────────────────
# API 엔드포인트
# ─────────────────────────────────────────

# ── 프론트엔드 서빙 ──
@app.get("/")
async def serve_index():
    return FileResponse("index.html")

@app.get("/health")
async def health():
    return {"status": "ok"}


# ── 녹음 제어 ──
@app.post("/recording/start")
async def recording_start():
    if session["is_recording"]:
        return {"status": "error", "message": "이미 녹음 중입니다."}
    session["is_recording"] = True
    session["transcript"]   = []
    session["start_time"]   = datetime.datetime.now()
    session["minutes_path"] = ""
    return {"status": "started"}


@app.post("/recording/chunk")
async def recording_chunk(
    file: UploadFile = File(...),
    language: str = "ko"
):
    """브라우저에서 보낸 오디오 청크를 Whisper로 전사"""
    if not session["is_recording"]:
        return {"status": "ignored"}

    audio_bytes = await file.read()
    if len(audio_bytes) < 1000:   # 너무 짧은 청크 무시
        return {"text": ""}

    suffix = "." + (file.filename.split(".")[-1] if file.filename else "webm")
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(audio_bytes)
        tmp_path = tmp.name

    try:
        with open(tmp_path, "rb") as f:
            result = client.audio.transcriptions.create(
                model="whisper-1", file=f, language=language)
        text = result.text.strip()
        if text and not _is_hallucination(text):
            session["transcript"].append(text)
            return {"text": text}
        return {"text": ""}
    except Exception as e:
        return {"text": "", "error": str(e)}
    finally:
        os.unlink(tmp_path)


@app.post("/recording/stop")
async def recording_stop():
    if not session["is_recording"]:
        return {"status": "error", "message": "녹음 중이 아닙니다."}

    session["is_recording"] = False
    fname = _generate_minutes()
    session["minutes_path"] = fname

    return {
        "status":       "stopped",
        "line_count":   len(session["transcript"]),
        "minutes_file": fname,
        "transcript":   session["transcript"],
    }


@app.get("/recording/status", response_model=RecordingStatus)
async def recording_status():
    return RecordingStatus(
        is_recording = session["is_recording"],
        line_count   = len(session["transcript"]),
        minutes_path = session["minutes_path"],
    )


@app.get("/download/{filename}")
async def download_minutes(filename: str):
    """생성된 회의록 파일 다운로드"""
    path = os.path.join(OUTPUT_DIR, filename)
    if not os.path.exists(path):
        return JSONResponse({"error": "파일을 찾을 수 없습니다."}, status_code=404)
    return FileResponse(
        path,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        filename=filename,
    )


# ── AI 채팅 ──
@app.post("/chat", response_model=ChatResponse)
async def chat(req: ChatRequest):
    messages = [
        {
            "role": "system",
            "content": (
                "너는 한국어를 사용하는 AI Assistant야. "
                "사용자의 요청을 파악해서 적절한 도구(tool)를 사용해. "
                "도구 실행 후 결과를 자연스럽게 한국어로 설명해줘. "
                "start_meeting_recorder 도구를 선택하면 result가 RECORDING_START_REQUESTED인데, "
                "이 경우 사용자에게 화면의 녹음 버튼을 눌러 시작하라고 안내해줘."
            )
        },
        {"role": "user", "content": req.message}
    ]

    tool_calls_result: list[ToolCallItem] = []

    response = client.chat.completions.create(
        model="gpt-4o",
        messages=messages,
        tools=TOOLS,
        tool_choice="auto",
    )
    msg = response.choices[0].message

    if msg.tool_calls:
        messages.append(msg)
        for tc in msg.tool_calls:
            fn_name = tc.function.name
            fn_args = json.loads(tc.function.arguments)
            display = TOOL_DISPLAY.get(fn_name, fn_name)
            try:
                result = _run_tool(fn_name, fn_args)
                status = "done"
            except Exception as e:
                result = f"오류: {e}"
                status = "error"

            tool_calls_result.append(ToolCallItem(
                name=display, description=result[:120], status=status))
            messages.append({"role":"tool","tool_call_id":tc.id,"content":result})

        final = client.chat.completions.create(model="gpt-4o", messages=messages)
        reply = final.choices[0].message.content.strip()
    else:
        reply = msg.content.strip()

    return ChatResponse(reply=reply, tool_calls=tool_calls_result)


# ─────────────────────────────────────────
# 실행
# ─────────────────────────────────────────
if __name__ == "__main__":
    import uvicorn
    if API_KEY == "여기에_API_KEY_입력":
        print("⚠️  OPENAI_API_KEY 환경변수를 설정해주세요.")
        exit(1)
    print(f"🚀 서버 시작: http://localhost:{PORT}")
    uvicorn.run(app, host="0.0.0.0", port=PORT)